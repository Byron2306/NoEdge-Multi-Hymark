#!/usr/bin/env python3
"""
Convert DOCX rubric files into HOMS internal JSON rubric format.

Best-effort extractor: prefers table-based rubrics, falls back to
paragraph heuristics. Writes outputs to `output/rubrics/<name>.internal.json`.
"""
import json
import re
from pathlib import Path

try:
    from docx import Document
except Exception:
    Document = None


def parse_table(table):
    rows = []
    for i, row in enumerate(table.rows):
        cells = [c.text.strip() for c in row.cells]
        # Skip empty rows
        if not any(cells):
            continue
        rows.append(cells)
    return rows


def rows_to_rubric(rows):
    rubric = {}
    for cells in rows:
        # Heuristic mapping: first column = criterion name
        name = cells[0]
        weight = 1.0
        annotations = []
        if len(cells) >= 2:
            # try parse weight from second cell
            wt_text = cells[1]
            m = re.search(r"([0-9]*\.?[0-9]+)", wt_text)
            if m:
                try:
                    weight = float(m.group(1))
                except Exception:
                    weight = 1.0
            else:
                # maybe annotations in second cell
                annotations = [a.strip() for a in re.split(r"[,;]|\n", wt_text) if a.strip()]
        if len(cells) >= 3:
            anns = [a.strip() for a in re.split(r"[,;]|\n", cells[2]) if a.strip()]
            annotations = (annotations or []) + anns

        rubric[name or 'Unnamed Criterion'] = {"weight": weight, "annotations": annotations}
    return rubric


def parse_paragraphs(doc):
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    rubric = {}
    for line in lines:
        # look for patterns like "1. Criterion name - weight: 1.0 - annotations: a,b"
        m = re.match(r"^\s*\d+\.\s*(.+)$", line)
        text = m.group(1) if m else line
        parts = re.split(r" - |:\s", text, maxsplit=2)
        if len(parts) == 1:
            # try split by ' — ' or ' – '
            parts = re.split(r" — | – | - ", text)
        name = parts[0].strip()
        weight = 1.0
        annotations = []
        if len(parts) >= 2:
            # find numbers in the rest
            m = re.search(r"([0-9]*\.?[0-9]+)", parts[1])
            if m:
                try:
                    weight = float(m.group(1))
                except Exception:
                    weight = 1.0
            else:
                annotations = [a.strip() for a in re.split(r"[,;]|\n", parts[1]) if a.strip()]
        if len(parts) >= 3:
            anns = [a.strip() for a in re.split(r"[,;]|\n", parts[2]) if a.strip()]
            annotations = (annotations or []) + anns

        rubric[name or 'Unnamed Criterion'] = {"weight": weight, "annotations": annotations}
    return rubric


def convert_file(path: Path, out_dir: Path):
    print(f"Converting {path}...")
    if Document is None:
        print("python-docx is not installed. Install with: pip install python-docx")
        return None
    doc = Document(path)
    rubric = {}
    # Prefer table-based extraction
    if doc.tables:
        all_rows = []
        for t in doc.tables:
            rows = parse_table(t)
            if rows:
                all_rows.extend(rows)
        if all_rows:
            rubric = rows_to_rubric(all_rows)
    # Fallback to paragraph parsing
    if not rubric:
        rubric = parse_paragraphs(doc)

    # If still empty, add full text as fallback
    if not rubric:
        full = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        rubric = {"FullText": {"weight": 1.0, "annotations": [full[:200]]}}

    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{path.stem}.internal.json"
    with out_path.open('w', encoding='utf8') as f:
        json.dump(rubric, f, indent=2, ensure_ascii=False)
    print(f"Wrote {out_path}")
    return out_path


def main():
    base = Path.cwd()
    candidates = [
        base / 'Rubric-1.docx',
        base / 'Rubric2.docx',
        base / 'Rubric_UTEW221_A1.docx',
        base / 'Rubric_UTEW221_A2.docx',
        base / 'UTEW221_Assignment2_Template_2025.docx'
    ]
    out_dir = base / 'output' / 'rubrics'
    found = [p for p in candidates if p.exists()]
    if not found:
        # also try scanning cwd for any docx
        found = list(Path('.').glob('*.docx'))
    if not found:
        print('No DOCX files found to convert in the workspace root.')
        return

    for p in found:
        try:
            convert_file(p, out_dir)
        except Exception as e:
            print(f"Failed to convert {p}: {e}")


if __name__ == '__main__':
    main()
